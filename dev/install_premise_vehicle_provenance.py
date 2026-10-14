"""Restore vehicle and legacy source metadata without changing any exchanges."""

import csv
import json
import shutil
from copy import deepcopy
from pathlib import Path

import openpyxl

from audit_premise_vehicle_workbooks import identity, read_workbook
from premise_vehicle_provenance import (
    KEY,
    attach_legacy_provenance,
    attach_provenance,
    digest,
    generation_manifest,
)

ROOT = Path(__file__).resolve().parents[1]
WORK = Path("/private/tmp/premise-vehicle-provenance-20261014")
OUT = ROOT / "results/premise_vehicle_provenance_20261014"
RAW = Path("/private/tmp/premise-vehicle-battery-simple-20261013")
INVENTORIES = ROOT.parent / "premise/premise/data/additional_inventories"
ADDED = {
    "source",
    "model",
    "model version",
    "model publication",
    KEY,
    "inventory provenance",
    "description",
    "special remark",
    "comment",
}


def normalized(value):
    if isinstance(value, dict):
        return {k: normalized(v) for k, v in value.items() if k != "row"}
    if isinstance(value, list):
        return [normalized(v) for v in value]
    return value


def patch_workbook(installed, candidate, before, expected):
    workbook = openpyxl.load_workbook(installed)
    for sheet in workbook:
        if str(sheet.cell(1, 1).value).lower() == "database":
            sheet.cell(1, 2, installed.stem)
    for ds in sorted(expected, key=lambda d: d["row"], reverse=True):
        sheet, row = workbook[ds["sheet"]], ds["row"] + 1
        seen = set()
        while str(sheet.cell(row, 1).value).lower() != "exchanges":
            field = str(sheet.cell(row, 1).value).lower()
            if field in ADDED and field in ds:
                sheet.cell(row, 2, ds[field])
                seen.add(field)
            row += 1
            assert row <= sheet.max_row
        missing = sorted(k for k in ADDED if k in ds and k not in seen)
        sheet.insert_rows(row, len(missing))
        for offset, field in enumerate(missing):
            sheet.cell(row + offset, 1, field)
            sheet.cell(row + offset, 2, ds[field])
    workbook.save(candidate)
    workbook.close()
    actual = read_workbook(candidate)[0]
    assert {identity(d): normalized(d) for d in actual} == {
        identity(d): normalized(d) for d in expected
    }
    assert {identity(d): normalized(d["exchanges"]) for d in actual} == {
        identity(d): normalized(d["exchanges"]) for d in before
    }


def main():
    if (OUT / "installation.json").exists():
        raise RuntimeError("Installation already recorded")
    (WORK / "candidates").mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    records, extra_files = [], []
    for family, folder, filename in (
        ("car", "cars", "lci-pass_cars.xlsx"),
        ("truck", "trucks", "lci-trucks.xlsx"),
        ("legacy", "trucks", "lci-truck-legacy-components.xlsx"),
    ):
        installed, candidate = INVENTORIES / filename, WORK / "candidates" / filename
        before = read_workbook(installed)[0]
        expected = deepcopy(before)
        assert not any(KEY in ds or "inventory provenance" in ds for ds in before)
        if family == "legacy":
            audit = ROOT / "results/premise_vehicle_audit_20261012"
            original = {
                identity(d): d
                for d in json.loads((audit / "parsed_inventories.json").read_text())
                if d["workbook"] == "lci-trucks.xlsx"
            }
            original_hash = json.loads((audit / "summary.json").read_text())["truck"][
                "sha256"
            ]
            attach_legacy_provenance(expected, original, original_hash)
        else:
            raw = RAW / folder
            record = json.loads((raw / "provenance.json").read_text())
            assert digest(raw / filename) == record["sha256"]
            # Later LPG metadata changed adapter code, not model calculations.
            for name, value in record["source_sha256"].items():
                file = Path(name) if Path(name).is_absolute() else ROOT.parent / name
                if "dev" not in file.parts:
                    assert digest(file) == value, name
            adapted = {
                identity(d): d for d in json.loads((raw / "adapted.json").read_text())
            }
            assert adapted.keys() == {identity(d) for d in expected}
            for ds in expected:
                inherited = adapted[identity(ds)]
                if inherited.get("source"):
                    ds["source"] = inherited["source"]
            parameters = list(csv.DictReader((raw / "vehicle_parameters.csv").open()))
            generation = generation_manifest(
                family, ROOT.parent, raw, record, "2026-10-13"
            )
            for artifact, suffix in (
                ("vehicle_parameters.csv", "parameters.csv"),
                ("noise_audit.csv", "noise.csv"),
            ):
                name = f"{installed.stem}.{suffix}"
                shutil.copy2(raw / artifact, WORK / "candidates" / name)
                generation["artifacts"][artifact]["companion file"] = name
                extra_files.append(name)
            attach_provenance(expected, family, generation, parameters)
            old_sidecar = installed.with_suffix(".provenance.json")
            sidecar = {
                "schema version": 2,
                "workbook": filename,
                "sha256": None,
                "generation provenance": generation,
                "previous provenance": json.loads(old_sidecar.read_text()),
                "metadata update": {
                    "date": "2026-10-14",
                    "before workbook sha256": digest(installed),
                    "before sidecar sha256": digest(old_sidecar),
                    "boundary": "Metadata and workbook database header only; no model regeneration or exchange changes.",
                },
            }
        for ds in expected:
            ds["database"] = installed.stem
            for field in ADDED:
                if isinstance(ds.get(field), (dict, list)):
                    ds[field] = json.dumps(ds[field], ensure_ascii=False)
                    assert len(ds[field]) < 32767
        patch_workbook(installed, candidate, before, expected)
        records.append(
            dict(
                filename=filename,
                before_sha256=digest(installed),
                after_sha256=digest(candidate),
                activities=len(expected),
                all_exchange_fields_unchanged=True,
                database_label=installed.stem,
            )
        )
        if family != "legacy":
            sidecar["sha256"] = digest(candidate)
            name = old_sidecar.name
            (WORK / "candidates" / name).write_text(
                json.dumps(sidecar, indent=2) + "\n"
            )
            extra_files.append(name)
    backup = WORK / "before"
    backup.mkdir(exist_ok=False)
    for name in [r["filename"] for r in records] + extra_files:
        installed = INVENTORIES / name
        if installed.exists():
            shutil.copy2(installed, backup / name)
        shutil.copy2(WORK / "candidates" / name, installed)
    manifest = dict(
        files=records,
        companions={
            name: {
                "after_sha256": digest(INVENTORIES / name),
                "before_sha256": (
                    digest(backup / name) if (backup / name).exists() else None
                ),
            }
            for name in extra_files
        },
        source_files={
            p.name: digest(p)
            for p in (
                Path(__file__),
                ROOT / "dev/premise_vehicle_provenance.py",
                ROOT / "dev/premise_vehicle_excel.py",
                ROOT / "dev/regenerate_premise_cars.py",
                ROOT / "dev/regenerate_premise_trucks.py",
            )
        },
    )
    (OUT / "installation.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps(records, indent=2))


if __name__ == "__main__":
    main()
