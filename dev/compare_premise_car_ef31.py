"""Compare the 16 EF 3.1 headline indicators in paired 2020 premise builds.

One vehicle-kilometre, exact service/product/location matches; no weighting or
normalization. Stores scores and foreground diagnostics, not full background LCI.

Requires the old/new build manifests in --output, parsed_inventories.json beside
the scope manifest, and adapted.json in --output's parent directory. Generate
these inputs with the workbook audit, car adapter, and paired build scripts.
"""

import argparse
import csv
import json
import os
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd


def format_workbook(path):
    """Keep the full result table filterable and small indicators readable."""
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        headers = {cell.column: cell.value for cell in sheet[1]}
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="187C80")
        for column, header in headers.items():
            letter = get_column_letter(column)
            width = {
                "name": 70,
                "category": 48,
                "method": 105,
                "functional_unit": 24,
                "activity_code": 38,
                "database": 60,
                "unit": 25,
            }.get(header, 19)
            sheet.column_dimensions[letter].width = width
            if header in {"old", "new", "absolute_change", "score"}:
                for cell in list(sheet.columns)[column - 1][1:]:
                    cell.number_format = "0.00000E+00"
            if header in {"percent_change", "min", "median", "max"}:
                for cell in list(sheet.columns)[column - 1][1:]:
                    cell.number_format = '0.0"%"'
            if header == "percent_change":
                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{sheet.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="A5D6A7",
                        mid_type="num",
                        mid_value=0,
                        mid_color="FFFFFF",
                        end_type="max",
                        end_color="EF9A9A",
                    ),
                )
    workbook.save(path)


def audit_template_emissions(diagnostic, original_path, updated_path, output):
    """Compare workbook emission totals with written EUR activities.

    Old ecoinvent 3.7 biosphere names can migrate, so old missing-name matches
    must not be interpreted solely as emission scaling.
    """
    rows = []
    for variant, path in (("old", original_path), ("new", updated_path)):
        index = {
            d["name"]: d
            for d in json.loads(path.read_text())
            if d["name"].startswith("transport, passenger car,")
        }
        subset = diagnostic[
            (diagnostic.variant == variant)
            & (diagnostic.location == "EUR")
            & (diagnostic.type == "biosphere")
        ]
        for (vehicle, flow), group in subset.groupby(["vehicle", "name"]):
            before = sum(
                e["amount"]
                for e in index[vehicle]["exchanges"]
                if e["type"] == "biosphere" and e["name"] == flow
            )
            after = group.amount.sum()
            rows.append(
                dict(
                    variant=variant,
                    vehicle=vehicle,
                    flow=flow,
                    template=before,
                    EUR=after,
                    ratio=after / before if before else None,
                    changed=not np.isclose(after, before, rtol=1e-5, atol=1e-14),
                )
            )
    pd.DataFrame(rows).to_csv(output / "template_to_EUR_emissions.csv", index=False)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--manifest",
        type=Path,
        default=Path("results/premise_vehicle_audit_20261012/candidate_recipes.csv"),
    )
    args = parser.parse_args()
    assert os.environ.get("BRIGHTWAY2_DIR"), "Use the isolated project copy"
    import bw2calc as bc
    import bw2data as bd

    bd.projects.set_current("ecoinvent-3.12-cutoff")
    args.output.mkdir(parents=True, exist_ok=True)
    recipes = [
        r
        for r in csv.DictReader(args.manifest.open())
        if r["family"] == "car" and r["selected_for_2020_update"] == "True"
    ]
    methods = sorted(
        m
        for m in bd.methods
        if m[1] == "EF v3.1"
        and not m[2].startswith("climate change:")
        and ", organics" not in m[2]
        and ", inorganics" not in m[2]
    )
    assert len(methods) == 16
    records, diagnostic, counts = [], [], {}
    method_metadata = [{"method": m, "unit": bd.methods[m]["unit"]} for m in methods]
    for variant in ("old", "new"):
        metadata = json.loads((args.output / f"{variant}_build.json").read_text())
        dbname = metadata["database"]
        index = {}
        for a in bd.Database(dbname):
            if a["location"] not in ("RER", "EUR"):
                continue
            key = (a["name"].lower(), a["reference product"].lower(), a["location"])
            index.setdefault(key, []).append(a)
        selected = []
        for recipe in recipes:
            for region in ("RER", "EUR"):
                matches = index.get(
                    (
                        recipe["name"].lower(),
                        recipe["reference product"].lower(),
                        region,
                    ),
                    [],
                )
                if region == "RER":
                    assert len(matches) == 1, (
                        variant,
                        recipe["name"],
                        region,
                        len(matches),
                    )
                elif not matches:
                    continue
                assert len(matches) == 1
                assert matches[0]["unit"] == "kilometer"
                selected.append((recipe, matches[0]))
        counts[variant] = dict(Counter(a["location"] for _, a in selected))
        print(variant, counts[variant], flush=True)
        lca = bc.LCA({selected[0][1]: 1}, methods[0])
        lca.lci(factorize=True)
        cf = []
        for method in methods:
            lca.switch_method(method)
            cf.append(np.asarray(lca.characterization_matrix.sum(axis=0)).ravel())
        cf = np.stack(cf)
        for i, (recipe, act) in enumerate(selected):
            lca.redo_lci({act.id: 1})
            flow_totals = lca.biosphere_matrix @ lca.supply_array
            scores = cf @ flow_totals
            assert np.isfinite(scores).all()
            for method, score in zip(methods, scores):
                records.append(
                    dict(
                        variant=variant,
                        name=recipe["name"],
                        product=act["reference product"],
                        location=act["location"],
                        size=recipe["model_size"] or "Medium",
                        powertrain=recipe["model_powertrain"] or "LPG proxy",
                        database=dbname,
                        activity_code=act["code"],
                        functional_unit="1 vehicle-kilometre",
                        category=method[2],
                        unit=bd.methods[method]["unit"],
                        score=float(score),
                    )
                )
            # Independently verify the vectorized characterization against Brightway.
            if i == 0:
                lca.switch_method(methods[0])
                lca.lcia()
                assert np.isclose(lca.score, scores[0], rtol=1e-10, atol=1e-14)
            for exc in act.exchanges():
                if exc["type"] == "production":
                    continue
                diagnostic.append(
                    dict(
                        variant=variant,
                        vehicle=recipe["name"],
                        location=act["location"],
                        type=exc["type"],
                        name=exc.input["name"],
                        supplier_location=exc.input.get("location", ""),
                        unit=exc.input["unit"],
                        amount=exc["amount"],
                        categories="::".join(exc.input.get("categories", ())),
                        supplier_database=exc.input["database"],
                    )
                )
        print("Scored", variant, len(selected), flush=True)
    assert counts["old"] == counts["new"]
    scores = pd.DataFrame(records)
    keys = [
        "name",
        "location",
        "size",
        "powertrain",
        "functional_unit",
        "category",
        "unit",
    ]
    paired = scores.pivot(index=keys, columns="variant", values="score").reset_index()
    assert not paired[["old", "new"]].isna().any().any()
    paired["absolute_change"] = paired["new"] - paired["old"]
    paired["percent_change"] = np.where(
        paired["old"] != 0, paired["absolute_change"] / paired["old"] * 100, np.nan
    )
    paired.to_csv(args.output / "ef31_comparison.csv", index=False)
    scores.to_csv(args.output / "ef31_scores.csv", index=False)
    diag = pd.DataFrame(diagnostic)
    diag.to_csv(args.output / "foreground_exchanges.csv", index=False)
    audit_template_emissions(
        diag,
        args.manifest.with_name("parsed_inventories.json"),
        args.output.parent / "adapted.json",
        args.output,
    )
    summary = (
        paired.groupby(["location", "category"])["percent_change"]
        .agg(["min", "median", "max"])
        .reset_index()
    )
    summary.to_csv(args.output / "ef31_summary.csv", index=False)
    with pd.ExcelWriter(
        args.output / "ef31_comparison.xlsx", engine="openpyxl"
    ) as writer:
        paired.to_excel(writer, sheet_name="Comparison", index=False)
        summary.to_excel(writer, sheet_name="Category summary", index=False)
        scores.to_excel(writer, sheet_name="Scores and activity keys", index=False)
        pd.DataFrame(method_metadata).astype(str).to_excel(
            writer, sheet_name="Methods", index=False
        )
    format_workbook(args.output / "ef31_comparison.xlsx")
    (args.output / "lcia_provenance.json").write_text(
        json.dumps(
            dict(
                project=bd.projects.current,
                brightway_directory=str(bd.projects.dir),
                functional_unit="1 vehicle-kilometre",
                methods=method_metadata,
                counts=counts,
                indicators=16,
                vectorized_characterization_verified=True,
                normalization=False,
                weighting=False,
            ),
            indent=2,
        )
        + "\n"
    )
    # Standalone scientific figure: absolute climate scores for every Medium car.
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    view = paired[
        (paired.category == "climate change")
        & (paired.location == "RER")
        & paired.name.str.contains(", Medium(?:,|$)", regex=True)
    ].copy()
    view = view.sort_values("old")
    labels = (
        view.name.str.replace("transport, passenger car, ", "", regex=False)
        .str.replace(", Medium", "", regex=False)
        .str.replace(", EURO-6", "", regex=False)
    )
    y = np.arange(len(view))
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.barh(
        y - 0.2, view.old * 1000, height=0.38, label="Old inventories", color="#9ca3af"
    )
    ax.barh(
        y + 0.2,
        view.new * 1000,
        height=0.38,
        label="Updated inventories",
        color="#187c80",
    )
    ax.set_yticks(y, labels)
    ax.set_xlabel("EF 3.1 climate change · g CO₂-eq per vehicle-km")
    ax.set_title(
        "Medium passenger cars · RER templates\nREMIND SSP2-NPi 2020 · ecoinvent 3.12 cutoff"
    )
    ax.legend(frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(args.output / "medium_cars_climate.png", dpi=180)
    fig.savefig(args.output / "medium_cars_climate.svg")
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()
