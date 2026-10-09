"""Independent source totals, temporal boundaries and hydrogen energy feedback."""

import importlib.util
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
import xarray as xr

from carculator_utils import DATA_DIR
from carculator_utils.electricity import get_electricity_mix, select_electricity_mix
from carculator_utils.hydrogen_power import (
    activity_key,
    fill_hydrogen_power,
    load_hydrogen_power,
    register_hydrogen_power,
)

EU27 = set(
    "AT BE BG HR CY CZ DK EE FI FR DE GR HU IE IT LV LT LU MT NL PL PT RO SK SI ES SE".split()
)
SCENARIO = "tyndp-2026-ntplus"


def test_country_coverage_endpoints_and_holds():
    mix = get_electricity_mix(SCENARIO)
    geco = get_electricity_mix()
    assert set(mix.country.values[mix.projection_last_year.values == 2050]) == EU27 | {
        "EU27"
    }
    for country in EU27:
        selected, provenance = select_electricity_mix(mix, country)
        assert provenance["projection_kind"] == "country"
        assert provenance["projection_region"] == country
        assert "NT+" in provenance["projection_source"]
        np.testing.assert_array_equal(selected.sel(year=2050), selected.sel(year=2070))
        np.testing.assert_allclose(
            selected.sel(year=2024, variable=geco.coords["variable"]),
            geco.sel(country=country, year=2024),
        )
    # Other countries retain GECO's actual national or disclosed regional source.
    np.testing.assert_allclose(
        mix.sel(country="US", variable=geco.coords["variable"]), geco.sel(country="US")
    )
    assert mix.sel(country="US").projection_last_year.item() == 2070
    _, provenance = select_electricity_mix(mix, "US")
    assert "GECO" in provenance["projection_source"]
    # Independent source output in TWh, rather than a duplicated mapping formula.
    balances = pd.read_csv(
        DATA_DIR / "electricity/tyndp_2026_balance_audit.csv", sep=";"
    ).set_index(["country", "year"])
    for country, nuclear in (("FR", 289.175668), ("PL", 76.017004)):
        total = balances.loc[(country, 2050), "mapped_generation_twh"]
        assert mix.sel(
            country=country, year=2050, variable="Nuclear"
        ).item() == pytest.approx(nuclear / total, abs=1e-8)
    assert mix.sel(country="DE", year=2050, variable="Nuclear").item() == 0
    np.testing.assert_allclose(
        balances.reported_generation_twh, balances.component_generation_twh, atol=1e-7
    )
    np.testing.assert_allclose(
        balances.reported_generation_twh - balances.pumped_storage_twh,
        balances.mapped_generation_twh,
        atol=1e-7,
    )


@pytest.fixture
def builder():
    spec = importlib.util.spec_from_file_location(
        "refresh_electricity",
        Path(__file__).resolve().parents[1] / "scripts/refresh_electricity.py",
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def source_fixture(tmp_path, builder):
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for country in ["EU27", *sorted(EU27)]:
        sheet = book.create_sheet(country)
        for col, year in zip((7, 12, 17, 22), (2030, 2035, 2040, 2050)):
            sheet.cell(3, col - 3, year)
            sheet.cell(4, col, "Weighted WS")
        rows = [("Generation [TWh]", None, None)]
        for label in [
            *builder.TYNDP_MAP,
            "Hydrogen GT",
            "Fuel Cell",
            "PS Turbine",
            "ENS",
            "RES Curtailment",
        ]:
            rows.append(
                (
                    None,
                    label,
                    {
                        "Nuclear": 60,
                        "Wind Onshore": 20,
                        "Hydrogen GT": 20,
                        "PS Turbine": 10,
                        "ENS": 3,
                        "RES Curtailment": 7,
                    }.get(label),
                )
            )
        rows += [
            ("Electricity - Flexibility", "Batteries Discharging", 5),
            ("H2", "SMR (Blue) and Pyrolisis", 2),
            (None, "SMR (Grey)", 1),
            (None, "Electrolyzers E-Market (P2G)", 7),
            (None, "Electrolyzers DRES (P2G)", None),
            (None, "Electrolyzers SRES (P2G)", None),
            (None, "H2 Adequacy Units", None),
            (None, "NH3 Imports - Low Band", 10),
            ("Heat", "Heat", 0),
            ("Energy Balance [TWh]", "Electricity Generation", 110),
            ("H2 +", "H2 Domestic Production", 10),
        ]
        for index, (section, label, value) in enumerate(rows, start=5):
            sheet.cell(index, 1, section)
            sheet.cell(index, 2, label)
            if value is not None:
                if country == "EU27" and (
                    label in builder.TYNDP_MAP
                    or label
                    in (
                        "Hydrogen GT",
                        "Fuel Cell",
                        "PS Turbine",
                        "ENS",
                        "RES Curtailment",
                        "Electricity Generation",
                    )
                ):
                    value *= 27
                for col in (7, 12, 17, 22):
                    sheet.cell(index, col, value)
    path = tmp_path / "nt.xlsx"
    book.save(path)
    return path


def test_import_balances_storage_and_unknown_hydrogen_supply(builder, tmp_path):
    projections, balances, hydrogen, _ = builder.read_tyndp(
        source_fixture(tmp_path, builder)
    )
    row = projections.iloc[0]
    assert row.Nuclear == pytest.approx(0.6)
    assert row.Wind == pytest.approx(0.2)
    # 20 TWh hydrogen power: 7/20 green, 11/20 grey (incl. imports), 2/20 blue.
    assert row["Hydrogen turbine, electrolysis"] == pytest.approx(0.07)
    assert row["Hydrogen turbine, reforming"] == pytest.approx(0.11)
    assert row["Hydrogen turbine, reforming CCS"] == pytest.approx(0.02)
    assert set(balances.mapped_generation_twh) == {100, 2700}
    assert hydrogen.imports_twh.unique().tolist() == [10]


def test_missing_output_cannot_be_silently_normalized(builder, tmp_path):
    import openpyxl

    path = source_fixture(tmp_path, builder)
    book = openpyxl.load_workbook(path)
    sheet = book["DE"]
    row = next(r for r in sheet if r[1].value == "Nuclear")
    row[6].value = None
    book.save(path)
    with pytest.raises(ValueError, match="Generation balance.*DE"):
        builder.read_tyndp(path)


def test_hydrogen_inventory_energy_feedback_and_residual_factors():
    config = load_hydrogen_power()
    keys = [activity_key(k) for k in config["fuels"].values()]
    keys += [
        activity_key(i["key"])
        for v in config["converters"].values()
        for i in v["inputs"]
    ]
    keys += [
        ("Hydrogen", ("air",), "kilogram"),
        ("grid", "RER", "kilowatt hour", "electricity"),
        (
            "electricity supply for fuel preparation",
            "DE",
            "kilowatt hour",
            "electricity, low voltage",
        ),
    ]
    inputs = {k: i for i, k in enumerate(keys)}
    mapping, pem = register_hydrogen_power(inputs, "DE")
    a = np.eye(len(inputs))[None, :, :, None]
    b = xr.DataArray(np.zeros((1, 1, len(inputs))))
    template = inputs[activity_key(config["fuels"]["electrolysis"])]
    a[:, inputs[keys[-2]], template, :] = -54
    b.values[:, :, template] = 0.123  # Nonzero residual must survive the copy.
    inv = SimpleNamespace(inputs=inputs, A=a, B=b, vm=SimpleNamespace(country="DE"))
    fill_hydrogen_power(inv, mapping, pem)
    prep = inputs[keys[-1]]
    gt = inputs[mapping["Hydrogen turbine, electrolysis"]]
    assert a[0, inputs[pem], gt, 0] == pytest.approx(-0.05 / 0.995)
    assert a[0, prep, inputs[pem], 0] == -54
    assert b.values[0, 0, inputs[pem]] == 0.123
    assert b.values[0, 0, gt] == 0  # Direct factors come from the explicit flows.
    assert (
        a[
            0,
            inputs[
                (
                    "Nitrogen oxides",
                    ("air", "non-urban air or from high stacks"),
                    "kilogram",
                )
            ],
            gt,
            0,
        ]
        == -0.000154
    )
    # A 10% H2 power share consumes upstream electricity; solve the physical loop.
    a[0, gt, prep, 0] = -0.1
    demand = np.zeros(len(inputs))
    demand[prep] = 1
    supply = np.linalg.solve(a[0, :, :, 0], demand)
    assert supply[prep] == pytest.approx(1 / (1 - 0.1 * 54 * 0.05 / 0.995))


@pytest.mark.family
@pytest.mark.parametrize("country", ["DE", "EU27"])
@pytest.mark.parametrize(
    "name,prefix,size,powertrain,kwargs",
    [
        ("carculator", "Car", "Medium", "BEV", {}),
        ("carculator_bus", "Bus", "13m-city", "BEV-depot", {}),
        ("carculator_truck", "Truck", "40t", "BEV", {"cycle": "Regional delivery"}),
        ("carculator_two_wheeler", "TwoWheeler", "Motorcycle 11-35kW", "BEV", {}),
    ],
)
def test_tyndp_runs_in_each_vehicle_family(
    name, prefix, size, powertrain, kwargs, country
):
    package = pytest.importorskip(name)
    ip = getattr(package, prefix + "InputParameters")()
    ip.static()
    _, array = package.fill_xarray_from_input_parameters(
        ip, scope={"size": [size], "powertrain": [powertrain], "year": [2040]}
    )
    # Isolate electricity supply from climate: an explicit illustrative bus profile.
    if prefix == "Bus":
        kwargs = {**kwargs, "ambient_temperature": 20.0}
    model = getattr(package, prefix + "Model")(array, country=country, **kwargs)
    model.set_all()
    before = model.array.copy(deep=True)
    inventory = getattr(package, "Inventory" + prefix)(
        model,
        scenario="static",
        background_configuration={"electricity scenario": SCENARIO},
    )
    assert inventory.electricity_mix.sizes["technology"] == 27
    assert inventory.electricity_provenance["projection_region"] == country
    assert np.isfinite(inventory.calculate_impacts()).all()
    xr.testing.assert_identical(model.array, before)


def test_eu27_has_its_own_history_and_generation_weighted_projections():
    geco = get_electricity_mix()
    assert geco.sel(
        country="EU27", year=2024, variable="Nuclear"
    ).item() == pytest.approx(649.550 / 2771.558)
    assert not np.allclose(
        geco.sel(country="EU27", year=2024), geco.sel(country="RER", year=2024)
    )
    _, provenance = select_electricity_mix(geco, "EU27")
    assert provenance["projection_kind"] == "aggregate"
    assert provenance["projection_region"] == "European Union"
    assert provenance["projection_last_year"] == 2070
    mix = get_electricity_mix(SCENARIO)
    _, provenance = select_electricity_mix(mix, "EU27")
    assert provenance["projection_kind"] == "aggregate"
    balance = pd.read_csv(
        DATA_DIR / "electricity/tyndp_2026_balance_audit.csv", sep=";"
    )
    for year in (2030, 2035, 2040, 2050):
        weights = (
            balance.loc[(balance.year == year) & balance.country.isin(EU27)]
            .set_index("country")
            .mapped_generation_twh
        )
        national = mix.sel(country=weights.index.tolist(), year=year)
        weighted = (national.values * weights.to_numpy()[:, None]).sum(
            axis=0
        ) / weights.sum()
        np.testing.assert_allclose(
            mix.sel(country="EU27", year=year), weighted, atol=1e-10
        )
        assert not np.allclose(
            mix.sel(country="EU27", year=year), national.mean("country")
        )


def test_inconsistent_eu27_source_aggregate_is_rejected(builder, tmp_path):
    import openpyxl

    path = source_fixture(tmp_path, builder)
    book = openpyxl.load_workbook(path)
    sheet = book["EU27"]
    for row in sheet:
        if row[1].value in ("Nuclear", "Electricity Generation"):
            row[6].value += 1
    book.save(path)
    with pytest.raises(ValueError, match="EU27 aggregate does not reconcile"):
        builder.read_tyndp(path)


@pytest.mark.family
@pytest.mark.export
def test_cyclic_supplies_are_scoped_and_exported_without_mutation(tmp_path):
    package = pytest.importorskip("carculator")
    pytest.importorskip("bw2io")
    ip = package.CarInputParameters()
    ip.static()
    _, array = package.fill_xarray_from_input_parameters(
        ip,
        scope={
            "size": ["Large", "Medium"],
            "powertrain": ["BEV"],
            "year": [2030, 2050],
        },
    )
    array.loc[dict(parameter="kilometers per year")] = 20000
    array.loc[dict(parameter="lifetime kilometers", size="Large")] = 100000
    array.loc[dict(parameter="lifetime kilometers", size="Medium")] = 400000
    model = package.CarModel(array, country="DE")
    model.set_all()
    config = {"electricity scenario": SCENARIO}
    inventory = package.InventoryCar(
        model, scenario="static", background_configuration=config
    )
    before = inventory.A.copy()
    result = inventory.calculate_impacts()
    assert inventory.electricity_supply_originals
    for size in ("Large", "Medium"):
        selected = deepcopy(model)
        selected.array = selected.array.sel(size=[size])
        individual = package.InventoryCar(
            selected, scenario="static", background_configuration=config
        )
        xr.testing.assert_allclose(
            individual.calculate_impacts(), result.sel(size=[size])
        )
    exported = inventory.export_lci(format="bw2io", directory=tmp_path)
    for importer in exported:
        generators = [
            a
            for a in importer.data
            if a["name"].startswith("electricity production, hydrogen")
        ]
        assert generators
        assert all(
            "TYNDP" in a["source"] and "efficiency" in a["comment"] for a in generators
        )
        assert all(
            any(e["type"] == "technosphere" and e["amount"] > 0 for e in a["exchanges"])
            for a in generators
        )
    assert all(
        SCENARIO in s
        for s in inventory.export_lci(
            software="simapro", format="string", directory=tmp_path
        )
    )
    np.testing.assert_array_equal(inventory.A, before)
    xr.testing.assert_identical(inventory.calculate_impacts(), result)
    # The established 21-column custom API still overrides this scenario.
    custom = np.zeros((2, 21))
    custom[:, 1] = 1
    fixed = package.InventoryCar(
        model,
        scenario="static",
        background_configuration={**config, "custom electricity mix": custom},
    )
    assert fixed.electricity_mix.sizes["technology"] == 21
    assert not fixed.hydrogen_power_activities
    np.testing.assert_array_equal(fixed.mix, custom)
